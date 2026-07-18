"""
文档管理服务 — Phase 2 OnlyOffice 集成

提供文档的 CRUD、版本管理、OnlyOffice 配置生成等功能。
"""

import os
import shutil
import uuid
from pathlib import Path
from typing import Optional

from sqlalchemy import select, func as sa_func
from sqlalchemy.ext.asyncio import AsyncSession
import structlog

from app.models.document import Document, DocumentVersion

logger = structlog.get_logger()

# 文档存储根目录
DOCUMENT_STORAGE_ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "data", "documents")

# 文件类型到 OnlyOffice 文档类型的映射
FILE_TYPE_MAP = {
    "docx": "word",
    "xlsx": "cell",
    "pptx": "slide",
}

# 模板目录
TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "data", "templates")


class DocumentService:
    """文档管理服务"""

    def __init__(self):
        self.storage_root = DOCUMENT_STORAGE_ROOT
        self.template_dir = TEMPLATE_DIR
        os.makedirs(self.storage_root, exist_ok=True)
        os.makedirs(self.template_dir, exist_ok=True)

    async def create_document(
        self,
        title: str,
        file_type: str,
        db: AsyncSession,
        template_id: Optional[str] = None,
    ) -> Document:
        """创建新文档

        Args:
            title: 文档标题
            file_type: 文件类型 (docx/xlsx/pptx)
            db: 数据库会话
            template_id: 可选的模板 ID，用于从模板创建

        Returns:
            创建的 Document 对象
        """
        if file_type not in FILE_TYPE_MAP:
            raise ValueError(f"不支持的文件类型: {file_type}，可选: {list(FILE_TYPE_MAP.keys())}")

        # 生成文件名和路径
        doc_id = str(uuid.uuid4())
        filename = f"{title}.{file_type}"
        doc_dir = os.path.join(self.storage_root, doc_id)
        os.makedirs(doc_dir, exist_ok=True)
        file_path = os.path.join(doc_dir, filename)

        # 如果指定了模板，从模板复制
        if template_id:
            template_path = os.path.join(self.template_dir, f"{template_id}.{file_type}")
            if os.path.exists(template_path):
                shutil.copy2(template_path, file_path)
                logger.info("从模板创建文档", template_id=template_id, doc_id=doc_id)
            else:
                logger.warning("模板不存在，创建空白文档", template_id=template_id)
                self._create_empty_file(file_path, file_type)
        else:
            self._create_empty_file(file_path, file_type)

        size_bytes = os.path.getsize(file_path) if os.path.exists(file_path) else 0

        document = Document(
            id=doc_id,
            title=title,
            filename=filename,
            file_path=file_path,
            file_type=file_type,
            size_bytes=size_bytes,
            version=1,
            metadata_json={},
        )
        db.add(document)

        # 创建初始版本记录
        version = DocumentVersion(
            document_id=doc_id,
            version_number=1,
            file_path=file_path,
            change_summary="初始版本",
        )
        db.add(version)

        await db.flush()
        await db.refresh(document)

        logger.info("文档已创建", doc_id=doc_id, title=title, file_type=file_type)
        return document

    async def get_document(self, doc_id: str, db: AsyncSession) -> Optional[Document]:
        """获取文档详情

        Args:
            doc_id: 文档 ID
            db: 数据库会话

        Returns:
            Document 对象，不存在返回 None
        """
        result = await db.execute(select(Document).where(Document.id == doc_id))
        return result.scalar_one_or_none()

    async def list_documents(
        self,
        db: AsyncSession,
        skip: int = 0,
        limit: int = 20,
        file_type: Optional[str] = None,
    ) -> tuple[list[Document], int]:
        """获取文档列表

        Args:
            db: 数据库会话
            skip: 跳过数量
            limit: 返回数量上限
            file_type: 可选的文件类型筛选

        Returns:
            (文档列表, 总数)
        """
        query = select(Document)
        if file_type:
            query = query.where(Document.file_type == file_type)

        # 计数
        count_query = select(sa_func.count()).select_from(query.subquery())
        total = (await db.execute(count_query)).scalar() or 0

        # 分页
        query = query.order_by(Document.updated_at.desc().nullsfirst()).offset(skip).limit(limit)
        result = await db.execute(query)
        documents = list(result.scalars().all())

        return documents, total

    async def update_document(
        self,
        doc_id: str,
        db: AsyncSession,
        title: Optional[str] = None,
    ) -> Optional[Document]:
        """更新文档元数据

        Args:
            doc_id: 文档 ID
            db: 数据库会话
            title: 新标题

        Returns:
            更新后的 Document 对象，不存在返回 None
        """
        document = await self.get_document(doc_id, db)
        if not document:
            return None

        if title is not None:
            document.title = title

        await db.flush()
        await db.refresh(document)
        logger.info("文档已更新", doc_id=doc_id)
        return document

    async def delete_document(self, doc_id: str, db: AsyncSession) -> bool:
        """删除文档

        Args:
            doc_id: 文档 ID
            db: 数据库会话

        Returns:
            是否删除成功
        """
        document = await self.get_document(doc_id, db)
        if not document:
            return False

        # 删除版本记录
        versions_result = await db.execute(
            select(DocumentVersion).where(DocumentVersion.document_id == doc_id)
        )
        for version in versions_result.scalars().all():
            await db.delete(version)

        # 删除文档记录
        await db.delete(document)

        # 删除文件
        doc_dir = os.path.join(self.storage_root, doc_id)
        if os.path.exists(doc_dir):
            shutil.rmtree(doc_dir, ignore_errors=True)

        logger.info("文档已删除", doc_id=doc_id)
        return True

    async def get_document_versions(self, doc_id: str, db: AsyncSession) -> list[DocumentVersion]:
        """获取文档的所有版本

        Args:
            doc_id: 文档 ID
            db: 数据库会话

        Returns:
            版本列表
        """
        result = await db.execute(
            select(DocumentVersion)
            .where(DocumentVersion.document_id == doc_id)
            .order_by(DocumentVersion.version_number.desc())
        )
        return list(result.scalars().all())

    async def save_document_version(
        self,
        doc_id: str,
        file_path: str,
        db: AsyncSession,
        change_summary: Optional[str] = None,
    ) -> DocumentVersion:
        """保存文档新版本

        Args:
            doc_id: 文档 ID
            file_path: 新版本文件路径
            db: 数据库会话
            change_summary: 变更说明

        Returns:
            新创建的 DocumentVersion 对象
        """
        document = await self.get_document(doc_id, db)
        if not document:
            raise ValueError(f"文档不存在: {doc_id}")

        # 获取当前最大版本号
        versions = await self.get_document_versions(doc_id, db)
        next_version = max((v.version_number for v in versions), default=0) + 1

        # 复制文件到版本目录
        version_dir = os.path.join(self.storage_root, doc_id, "versions")
        os.makedirs(version_dir, exist_ok=True)
        version_filename = f"v{next_version}_{document.filename}"
        version_path = os.path.join(version_dir, version_filename)
        if os.path.exists(file_path):
            shutil.copy2(file_path, version_path)

        # 更新文档主记录
        document.version = next_version
        document.file_path = file_path
        if os.path.exists(file_path):
            document.size_bytes = os.path.getsize(file_path)

        # 创建版本记录
        version = DocumentVersion(
            document_id=doc_id,
            version_number=next_version,
            file_path=version_path,
            change_summary=change_summary or f"版本 {next_version}",
        )
        db.add(version)
        await db.flush()
        await db.refresh(version)

        logger.info("文档版本已保存", doc_id=doc_id, version=next_version)
        return version

    def get_onlyoffice_config(self, doc_id: str, document: Document) -> dict:
        """获取 OnlyOffice 编辑器配置

        Args:
            doc_id: 文档 ID
            document: Document 对象

        Returns:
            OnlyOffice 编辑器配置字典
        """
        from app.services.onlyoffice_service import OnlyOfficeService
        oo_service = OnlyOfficeService()
        return oo_service.generate_editor_config(doc_id, mode="edit")

    def _create_empty_file(self, file_path: str, file_type: str) -> None:
        """创建空白文档文件

        Args:
            file_path: 文件路径
            file_type: 文件类型
        """
        if file_type == "docx":
            self._create_empty_docx(file_path)
        elif file_type == "xlsx":
            self._create_empty_xlsx(file_path)
        elif file_type == "pptx":
            self._create_empty_pptx(file_path)

    def _create_empty_docx(self, file_path: str) -> None:
        """创建空白 docx 文件"""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument()
            doc.save(file_path)
        except ImportError:
            # python-docx 不可用时，写入最小有效 docx
            import zipfile
            with zipfile.ZipFile(file_path, 'w') as zf:
                zf.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
                zf.writestr('word/_rels/document.xml.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
                zf.writestr('word/document.xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t/></w:r></w:p></w:body></w:document>')
                zf.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')

    def _create_empty_xlsx(self, file_path: str) -> None:
        """创建空白 xlsx 文件"""
        try:
            from openpyxl import Workbook
            wb = Workbook()
            wb.save(file_path)
        except ImportError:
            import zipfile
            with zipfile.ZipFile(file_path, 'w') as zf:
                zf.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
                zf.writestr('xl/workbook.xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>')
                zf.writestr('xl/_rels/workbook.xml.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
                zf.writestr('xl/worksheets/sheet1.xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>')
                zf.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')

    def _create_empty_pptx(self, file_path: str) -> None:
        """创建空白 pptx 文件"""
        try:
            from pptx import Presentation
            prs = Presentation()
            prs.save(file_path)
        except ImportError:
            import zipfile
            with zipfile.ZipFile(file_path, 'w') as zf:
                zf.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/><Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/></Types>')
                zf.writestr('ppt/presentation.xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>')
                zf.writestr('ppt/_rels/presentation.xml.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/></Relationships>')
                zf.writestr('ppt/slides/slide1.xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:spTree></p:cSld></p:sld>')
                zf.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/></Relationships>')

    def list_templates(self) -> list[dict]:
        """列出可用的文档模板

        Returns:
            模板列表
        """
        templates = []
        if not os.path.exists(self.template_dir):
            return templates

        for fname in sorted(os.listdir(self.template_dir)):
            fpath = os.path.join(self.template_dir, fname)
            if os.path.isfile(fpath):
                name, ext = os.path.splitext(fname)
                ext = ext.lstrip(".")
                if ext in FILE_TYPE_MAP:
                    templates.append({
                        "id": name,
                        "filename": fname,
                        "file_type": ext,
                        "size_bytes": os.path.getsize(fpath),
                    })
        return templates


# 全局单例
_document_service: Optional[DocumentService] = None


def get_document_service() -> DocumentService:
    """获取文档服务单例"""
    global _document_service
    if _document_service is None:
        _document_service = DocumentService()
    return _document_service
